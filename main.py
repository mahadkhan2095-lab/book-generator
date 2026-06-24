import os
import base64
from dotenv import load_dotenv
load_dotenv()
from groq import Groq
from supabase import create_client
from docx import Document
from docx.shared import Pt
import resend

client = Groq()
supabase = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_KEY"])
resend.api_key = os.environ["RESEND_API_KEY"]
MODEL = "llama-3.3-70b-versatile"


# ── Outline ────────────────────────────────────────────────────────────────────

def generate_outline(title: str, notes: str) -> str:
    prompt = (
        f"Generate a detailed book outline for the following title.\n"
        f"Title: {title}\n"
        f"Instructions: {notes}"
    )
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}]
    )
    return response.choices[0].message.content


# ── Chapters ───────────────────────────────────────────────────────────────────

def generate_chapter(title: str, outline: str, chapter_title: str, previous_summaries: list[str]) -> str:
    summaries_text = "\n".join(previous_summaries) if previous_summaries else "None"
    prompt = (
        f"You are writing a book titled '{title}'.\n"
        f"Here is the full outline:\n{outline}\n\n"
        f"Using the following previous chapter summaries:\n{summaries_text}\n\n"
        f"Write the full content for this chapter: {chapter_title}\n"
        f"Be detailed, engaging, and well-structured."
    )
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}]
    )
    return response.choices[0].message.content


def generate_summary(chapter_content: str) -> str:
    prompt = (
        f"Write a short summary (3-5 lines) of the following chapter content:\n\n"
        f"{chapter_content}"
    )
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}]
    )
    return response.choices[0].message.content


def parse_chapter_titles(outline: str) -> list[str]:
    prompt = (
        f"Extract only the chapter titles from the following book outline.\n"
        f"Return them as a plain numbered list, one per line, with no extra text.\n\n"
        f"{outline}"
    )
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}]
    )
    raw = response.choices[0].message.content
    chapters = []
    for line in raw.strip().splitlines():
        line = line.strip()
        if line:
            cleaned = line.lstrip("0123456789.).  ").strip()
            if cleaned:
                chapters.append(cleaned)
    return chapters


# ── Supabase I/O ───────────────────────────────────────────────────────────────

def load_book(title: str) -> dict | None:
    result = supabase.table("books").select("*").eq("title", title).order("created_at").limit(1).execute()
    return result.data[0] if result.data else None


def save_book(title: str, notes: str, outline: str) -> dict:
    result = supabase.table("books").insert({
        "title": title,
        "notes_on_outline_before": notes,
        "outline": outline,
        "status_outline_notes": "pending",
        "final_review_notes_status": "pending"
    }).execute()
    return result.data[0]


def update_book(book_id: str, fields: dict):
    supabase.table("books").update(fields).eq("id", book_id).execute()


def load_chapters(book_id: str) -> list:
    result = supabase.table("chapters").select("*").eq("book_id", book_id).order("chapter_number").execute()
    return result.data


def save_chapters(book_id: str, chapter_titles: list[str]):
    rows = [
        {
            "book_id": book_id,
            "chapter_number": i + 1,
            "title": ch_title,
            "content": "",
            "summary": "",
            "chapter_notes_status": "pending",
            "chapter_notes": ""
        }
        for i, ch_title in enumerate(chapter_titles)
    ]
    supabase.table("chapters").insert(rows).execute()


def update_chapter(chapter_id: str, fields: dict):
    supabase.table("chapters").update(fields).eq("id", chapter_id).execute()


# ── Compilation & Export ───────────────────────────────────────────────────────

def compile_book(book: dict, chapters: list) -> str:
    lines = []
    lines.append(f"Title: {book['title']}\n")
    lines.append("=" * 50)
    lines.append("=== Outline ===")
    lines.append("=" * 50)
    lines.append(book["outline"])
    lines.append("")
    for chapter in chapters:
        lines.append("=" * 50)
        lines.append(f"=== Chapter {chapter['chapter_number']}: {chapter['title']} ===")
        lines.append("=" * 50)
        lines.append(chapter["content"])
        lines.append("")
    return "\n".join(lines)


def export_to_txt(compiled_text: str):
    with open("final_book.txt", "w", encoding="utf-8") as f:
        f.write(compiled_text)
    print("Book successfully compiled and saved as final_book.txt")


def export_to_docx(book: dict, chapters: list):
    doc = Document()

    # Title
    title_heading = doc.add_heading(book["title"], level=0)
    title_heading.runs[0].font.size = Pt(24)

    # Outline
    doc.add_heading("Outline", level=1)
    doc.add_paragraph(book["outline"])
    doc.add_page_break()

    # Chapters
    for chapter in chapters:
        doc.add_heading(f"Chapter {chapter['chapter_number']}: {chapter['title']}", level=1)
        doc.add_paragraph(chapter["content"])
        doc.add_page_break()

    doc.save("final_book.docx")
    print("Book successfully exported as final_book.docx")


def send_email(receiver_email: str):
    with open("final_book.docx", "rb") as f:
        file_content = base64.b64encode(f.read()).decode("utf-8")

    params: resend.Emails.SendParams = {
        "from": "onboarding@resend.dev",
        "to": [receiver_email],
        "subject": "Your Generated Book is Ready",
        "text": "Hi,\n\nYour book has been successfully generated. Please find the attached DOCX file.\n\nEnjoy your book!",
        "attachments": [{"filename": "final_book.docx", "content": file_content}]
    }
    resend.Emails.send(params)
    print(f"Book emailed successfully to {receiver_email}")


# ── Chapter Generation Loop ────────────────────────────────────────────────────

def run_chapter_generation(book: dict):
    title = book["title"]
    outline = book["outline"]
    book_id = book["id"]

    chapters = load_chapters(book_id)

    if not chapters:
        print("Parsing chapter titles from outline...\n")
        chapter_titles = parse_chapter_titles(outline)
        save_chapters(book_id, chapter_titles)
        chapters = load_chapters(book_id)
        print(f"Found {len(chapters)} chapters. Starting generation...\n")

    for i, chapter in enumerate(chapters):
        num = chapter["chapter_number"]
        ch_title = chapter["title"]
        status = chapter["chapter_notes_status"]

        print(f"\n{'='*50}")
        print(f"Chapter {num}: {ch_title}")
        print(f"Status: {status}")
        print('='*50)

        if status == "approved" and chapter["content"]:
            print("Already approved. Skipping.")
            continue

        elif status == "pending" and chapter["content"]:
            print("Waiting for chapter review...")
            break

        elif status == "needs_changes":
            notes = chapter.get("chapter_notes", "")
            print(f"Regenerating with notes: {notes}\n")
            previous_summaries = [
                f"Chapter {chapters[j]['chapter_number']} - {chapters[j]['title']}: {chapters[j]['summary']}"
                for j in range(i) if chapters[j]["summary"]
            ]
            content = generate_chapter(title, outline, ch_title, previous_summaries)
            summary = generate_summary(content)
            update_chapter(chapter["id"], {
                "content": content,
                "summary": summary,
                "chapter_notes_status": "pending"
            })
            print(content)
            print("\nRegenerated. Status reset to 'pending'. Review again.")
            break

        else:
            previous_summaries = [
                f"Chapter {chapters[j]['chapter_number']} - {chapters[j]['title']}: {chapters[j]['summary']}"
                for j in range(i) if chapters[j]["summary"]
            ]
            print("Generating chapter content...\n")
            content = generate_chapter(title, outline, ch_title, previous_summaries)
            summary = generate_summary(content)
            update_chapter(chapter["id"], {
                "content": content,
                "summary": summary,
                "chapter_notes_status": "pending"
            })
            print(content)
            print(f"\n--- Summary ---\n{summary}")
            print("\nStatus set to 'pending'. Review before continuing.")
            break


# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    title = "The Art of Deep Focus"
    notes_on_outline_before = "Focus on productivity, mindfulness, and overcoming digital distractions."

    if not notes_on_outline_before.strip():
        print("Notes missing. Stopping.")
        return

    book = load_book(title)

    if book:
        print(f"Found existing book in Supabase. Loading...\n")
    else:
        print("Generating outline...\n")
        outline = generate_outline(title, notes_on_outline_before)
        book = save_book(title, notes_on_outline_before, outline)
        print("Outline saved to Supabase.\n")

    print("=== Book Outline ===\n")
    print(book["outline"])

    status = book.get("status_outline_notes", "pending")
    print(f"\n--- Outline Status: {status} ---")

    if status == "pending":
        print("Waiting for review...")

    elif status == "approved":
        print("Outline approved. Starting chapter generation...\n")
        run_chapter_generation(book)

        chapters = load_chapters(book["id"])
        all_approved = all(ch["chapter_notes_status"] == "approved" for ch in chapters)

        if all_approved:
            if not book.get("final_review_notes_status"):
                update_book(book["id"], {"final_review_notes_status": "pending"})
                book["final_review_notes_status"] = "pending"

            final_status = book["final_review_notes_status"]
            print(f"\n--- Final Review Status: {final_status} ---")

            if final_status == "pending":
                print("Waiting for final review...")

            elif final_status == "approved":
                print("Final review approved. Compiling book...\n")
                compiled = compile_book(book, chapters)
                export_to_txt(compiled)
                export_to_docx(book, chapters)
                send_email("mahadkhan2095@gmail.com")

            elif final_status == "needs_changes":
                print("Book needs changes. Please update the relevant chapters and re-run.")

    elif status == "needs_changes":
        updated_notes = book.get("notes_on_outline_before", notes_on_outline_before)
        print("Regenerating outline with updated notes...\n")
        new_outline = generate_outline(book["title"], updated_notes)
        update_book(book["id"], {"outline": new_outline, "status_outline_notes": "pending"})
        print("=== Regenerated Outline ===\n")
        print(new_outline)
        print("\nStatus reset to 'pending'. Review again.")


if __name__ == "__main__":
    main()
