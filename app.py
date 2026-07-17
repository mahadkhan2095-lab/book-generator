from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.units import inch
import os
import base64
from dotenv import load_dotenv
load_dotenv()

from fastapi import FastAPI, Request, Form
from fastapi.responses import HTMLResponse, RedirectResponse, FileResponse, Response
from fastapi.templating import Jinja2Templates
from groq import Groq
from supabase import create_client
from docx import Document
from docx.shared import Pt
import resend

app = FastAPI()
templates = Jinja2Templates(directory="templates")

client = Groq()
supabase = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_KEY"])
resend.api_key = os.environ["RESEND_API_KEY"]
MODEL = "llama-3.3-70b-versatile"


# ── LLM Functions ──────────────────────────────────────────────────────────────

def generate_outline(title: str, notes: str) -> str:
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": f"Generate a detailed book outline.\nTitle: {title}\nInstructions: {notes}"}]
    )
    return response.choices[0].message.content


def parse_chapter_titles(outline: str) -> list[str]:
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": f"Extract only chapter titles from this outline as a plain numbered list, one per line, no extra text.\n\n{outline}"}]
    )
    raw = response.choices[0].message.content
    chapters = []
    for line in raw.strip().splitlines():
        cleaned = line.strip().lstrip("0123456789.).  ").strip()
        if cleaned:
            chapters.append(cleaned)
    return chapters


def generate_chapter(title: str, outline: str, chapter_title: str, previous_summaries: list[str]) -> str:
    summaries_text = "\n".join(previous_summaries) if previous_summaries else "None"
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": f"You are writing a book titled '{title}'.\nOutline:\n{outline}\n\nPrevious chapter summaries:\n{summaries_text}\n\nWrite full content for: {chapter_title}\nBe detailed, engaging, and well-structured."}]
    )
    return response.choices[0].message.content


def generate_summary(content: str) -> str:
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": f"Write a 3-5 line summary of this chapter:\n\n{content}"}]
    )
    return response.choices[0].message.content


# ── Supabase Functions ─────────────────────────────────────────────────────────

def load_book(title: str):
    result = supabase.table("books").select("*").eq("title", title).order("created_at").limit(1).execute()
    return result.data[0] if result.data else None

def load_book_by_id(book_id: str):
    result = supabase.table("books").select("*").eq("id", book_id).execute()
    return result.data[0] if result.data else None

def load_all_books():
    result = supabase.table("books").select("*").order("created_at", desc=True).execute()
    return result.data

def save_book(title: str, notes: str, outline: str):
    result = supabase.table("books").insert({
        "title": title, "notes_on_outline_before": notes,
        "outline": outline, "status_outline_notes": "pending",
        "final_review_notes_status": "pending"
    }).execute()
    return result.data[0]

def update_book(book_id: str, fields: dict):
    supabase.table("books").update(fields).eq("id", book_id).execute()

def load_chapters(book_id: str):
    result = supabase.table("chapters").select("*").eq("book_id", book_id).order("chapter_number").execute()
    return result.data

def save_chapters(book_id: str, titles: list[str]):
    rows = [{"book_id": book_id, "chapter_number": i+1, "title": t, "content": "", "summary": "", "chapter_notes_status": "pending", "chapter_notes": ""} for i, t in enumerate(titles)]
    supabase.table("chapters").insert(rows).execute()

def update_chapter(chapter_id: str, fields: dict):
    supabase.table("chapters").update(fields).eq("id", chapter_id).execute()


# ── Export Functions ───────────────────────────────────────────────────────────

def export_to_docx(book: dict, chapters: list):
    doc = Document()
    title_heading = doc.add_heading(book["title"], level=0)
    title_heading.runs[0].font.size = Pt(24)
    doc.add_heading("Outline", level=1)
    doc.add_paragraph(book["outline"])
    doc.add_page_break()
    for ch in chapters:
        doc.add_heading(f"Chapter {ch['chapter_number']}: {ch['title']}", level=1)
        doc.add_paragraph(ch["content"])
        doc.add_page_break()
    doc.save("final_book.docx")


def export_to_pdf(book: dict, chapters: list):
    doc = SimpleDocTemplate("final_book.pdf", pagesize=A4,
                            rightMargin=inch, leftMargin=inch,
                            topMargin=inch, bottomMargin=inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Title"], fontSize=24, spaceAfter=20)
    heading_style = ParagraphStyle("Heading", parent=styles["Heading1"], fontSize=16, spaceAfter=12)
    body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=11, spaceAfter=8, leading=16)

    story = []
    story.append(Paragraph(book["title"], title_style))
    story.append(Spacer(1, 0.3 * inch))
    story.append(Paragraph("Outline", heading_style))
    for line in book["outline"].split("\n"):
        if line.strip():
            story.append(Paragraph(line.strip(), body_style))
    story.append(PageBreak())

    for ch in chapters:
        story.append(Paragraph(f"Chapter {ch['chapter_number']}: {ch['title']}", heading_style))
        story.append(Spacer(1, 0.15 * inch))
        for line in ch["content"].split("\n"):
            if line.strip():
                story.append(Paragraph(line.strip(), body_style))
        story.append(PageBreak())

    doc.build(story)


def send_email(receiver_email: str):
    with open("final_book.docx", "rb") as f:
        file_content = base64.b64encode(f.read()).decode("utf-8")
    params: resend.Emails.SendParams = {
        "from": "onboarding@resend.dev",
        "to": [receiver_email],
        "subject": "Your Generated Book is Ready",
        "text": "Hi,\n\nYour book has been successfully generated. Please find the attached DOCX file.\n\nEnjoy your book!",
        "attachments": [{"filename": "final_book.docx", "content": file_content}]
    }
    resend.Emails.send(params)


# ── Routes ─────────────────────────────────────────────────────────────────────

@app.get("/favicon.ico", include_in_schema=False)
async def favicon():
    return Response(status_code=204)


@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    try:
        books = load_all_books()
    except Exception as e:
        print(f"Error loading books: {e}")
        books = []
    return templates.TemplateResponse(request, "index.html", {"books": books})


@app.post("/create-book")
async def create_book(title: str = Form(...), notes: str = Form(...)):
    existing = load_book(title)
    if existing:
        return RedirectResponse(f"/book/{existing['id']}", status_code=302)
    outline = generate_outline(title, notes)
    book = save_book(title, notes, outline)
    return RedirectResponse(f"/book/{book['id']}", status_code=302)


@app.get("/book/{book_id}", response_class=HTMLResponse)
async def book_dashboard(request: Request, book_id: str):
    book = load_book_by_id(book_id)
    chapters = load_chapters(book_id)
    return templates.TemplateResponse(request, "book.html", {"book": book, "chapters": chapters})


@app.post("/book/{book_id}/approve-outline")
async def approve_outline(book_id: str):
    update_book(book_id, {"status_outline_notes": "approved"})
    # Parse and save chapters
    book = load_book_by_id(book_id)
    chapters = load_chapters(book_id)
    if not chapters:
        titles = parse_chapter_titles(book["outline"])
        save_chapters(book_id, titles)
    return RedirectResponse(f"/book/{book_id}", status_code=302)


@app.post("/book/{book_id}/generate-chapter")
async def generate_next_chapter(book_id: str):
    book = load_book_by_id(book_id)
    chapters = load_chapters(book_id)
    for i, chapter in enumerate(chapters):
        if chapter["chapter_notes_status"] == "pending" and not chapter["content"]:
            previous_summaries = [
                f"Chapter {chapters[j]['chapter_number']} - {chapters[j]['title']}: {chapters[j]['summary']}"
                for j in range(i) if chapters[j]["summary"]
            ]
            content = generate_chapter(book["title"], book["outline"], chapter["title"], previous_summaries)
            summary = generate_summary(content)
            update_chapter(chapter["id"], {"content": content, "summary": summary})
            break
    return RedirectResponse(f"/book/{book_id}", status_code=302)


@app.post("/book/{book_id}/approve-all-chapters")
async def approve_all_chapters(book_id: str):
    chapters = load_chapters(book_id)
    for chapter in chapters:
        if chapter["chapter_notes_status"] != "approved" and chapter["content"]:
            update_chapter(chapter["id"], {"chapter_notes_status": "approved"})
    return RedirectResponse(f"/book/{book_id}", status_code=302)


@app.post("/chapter/{chapter_id}/approve")
async def approve_chapter(chapter_id: str, book_id: str = Form(...)):
    update_chapter(chapter_id, {"chapter_notes_status": "approved"})
    return RedirectResponse(f"/book/{book_id}", status_code=302)


@app.post("/chapter/{chapter_id}/needs-changes")
async def chapter_needs_changes(chapter_id: str, book_id: str = Form(...), notes: str = Form(...)):
    update_chapter(chapter_id, {"chapter_notes_status": "needs_changes", "chapter_notes": notes})
    return RedirectResponse(f"/book/{book_id}", status_code=302)


@app.post("/book/{book_id}/compile")
async def compile_book(book_id: str, email: str = Form(...)):
    book = load_book_by_id(book_id)
    chapters = load_chapters(book_id)
    export_to_docx(book, chapters)
    export_to_pdf(book, chapters)
    send_email("mahadkhan2095@gmail.com")
    update_book(book_id, {"final_review_notes_status": "approved"})
    return RedirectResponse(f"/book/{book_id}?compiled=true", status_code=302)


@app.get("/book/{book_id}/download-pdf")
async def download_pdf(book_id: str):
    return FileResponse("final_book.pdf", filename="final_book.pdf", media_type="application/pdf")


@app.get("/book/{book_id}/download")
async def download_book(book_id: str):
    return FileResponse("final_book.docx", filename="final_book.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
